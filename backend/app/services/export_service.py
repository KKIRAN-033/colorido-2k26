"""Export service - CSV, Excel, and PDF generation for registrations."""

import io
import csv
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


EXPORT_COLUMNS = [
    ("registration_id", "Registration ID"),
    ("participant_name", "Name"),
    ("college", "College"),
    ("event_name", "Event"),
    ("event_category", "Category"),
    ("event_division", "Division"),
    ("phone", "Phone"),
    ("email", "Email"),
    ("team_name", "Team Name"),
    ("participation_type", "Type"),
    ("status", "Status"),
    ("created_at", "Registered On"),
]


def format_value(value):
    """Format a value for export."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return str(value)


def generate_csv(registrations: List[dict]) -> io.StringIO:
    """Generate CSV file from registration data."""
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header row
    writer.writerow([col[1] for col in EXPORT_COLUMNS])
    
    # Data rows
    for reg in registrations:
        row = [format_value(reg.get(col[0], "")) for col in EXPORT_COLUMNS]
        writer.writerow(row)
    
    output.seek(0)
    return output


def generate_excel(registrations: List[dict]) -> io.BytesIO:
    """Generate Excel XLSX file from registration data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    
    # Styles
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=11)
    cell_alignment = Alignment(vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    
    # Title row
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = "COLORIDO 2K26 — Registration Data"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="e94560")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40
    
    # Date row
    ws.merge_cells("A2:L2")
    date_cell = ws["A2"]
    date_cell.value = f"Exported on: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    date_cell.font = Font(name="Calibri", size=10, italic=True)
    date_cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 25
    
    # Header row
    header_row = 4
    for col_idx, (_, col_name) in enumerate(EXPORT_COLUMNS, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 30
    
    # Data rows
    for row_idx, reg in enumerate(registrations, header_row + 1):
        for col_idx, (col_key, _) in enumerate(EXPORT_COLUMNS, 1):
            value = format_value(reg.get(col_key, ""))
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border
        
        # Alternate row colors
        if row_idx % 2 == 0:
            for col_idx in range(1, len(EXPORT_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    start_color="f0f0f5", end_color="f0f0f5", fill_type="solid"
                )
    
    # Column widths
    column_widths = [18, 22, 25, 22, 12, 12, 15, 25, 20, 10, 12, 18]
    for idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_registration_pdf(registration: dict) -> io.BytesIO:
    """Generate a PDF for a single registration pass with official formatting."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=1.5*cm,
        bottomMargin=1.5*cm,
        leftMargin=1.8*cm,
        rightMargin=1.8*cm,
    )
    styles = getSampleStyleSheet()
    elements = []
    
    # Title - COLORIDO 2K26
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=24,
        leading=28,
        textColor=colors.HexColor("#e94560"),
        alignment=1,
        fontName="Helvetica-Bold",
        spaceAfter=4,
    )
    elements.append(Paragraph("COLORIDO 2K26", title_style))
    
    # Subtitle - Credential Pass
    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["Normal"],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0f172a"),
        alignment=1,
        fontName="Helvetica-Bold",
        spaceAfter=4,
    )
    elements.append(Paragraph("OFFICIAL CREDENTIAL ENTRY PASS", subtitle_style))

    subtag_style = ParagraphStyle(
        "Subtag",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#64748b"),
        alignment=1,
        spaceAfter=16,
    )
    elements.append(Paragraph("National-Level College Cultural & Sports Festival • Multiverse Arena", subtag_style))
    elements.append(Spacer(1, 10))
    
    # Registration ID Highlight Banner
    reg_id = registration.get("registration_id", "N/A")
    banner_data = [[
        Paragraph(f"<b>REGISTRATION ID:</b> <font color='#e94560' size='+2'><b>{reg_id}</b></font>", ParagraphStyle("BannerLeft", parent=styles["Normal"], fontName="Helvetica", fontSize=11, textColor=colors.HexColor("#0f172a"))),
        Paragraph("<font color='#059669'><b>STATUS: CONFIRMED</b></font><br/><font size='8' color='#64748b'>100% Free Entry Verified</font>", ParagraphStyle("BannerRight", parent=styles["Normal"], alignment=2, fontName="Helvetica", fontSize=10))
    ]]
    banner_table = Table(banner_data, colWidths=[10*cm, 7.4*cm])
    banner_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fef2f2")),
        ("BOX", (0, 0), (-1, -1), 1.5, colors.HexColor("#f87171")),
        ("PADDING", (0, 0), (-1, -1), 10),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(banner_table)
    elements.append(Spacer(1, 15))

    # Registration details table
    data = [
        ["Participant / Captain", registration.get("participant_name", "N/A")],
        ["College / Institution", registration.get("college", "N/A")],
        ["Registered Event", registration.get("event_name", "N/A")],
        ["Arena Category", (registration.get("event_category") or "Cultural").capitalize()],
        ["Division / Format", registration.get("event_division") or "Open"],
        ["Participation Type", (registration.get("participation_type") or "Solo").capitalize()],
        ["Email Address", registration.get("email", "N/A")],
        ["Contact Phone", registration.get("phone", "N/A")],
    ]

    if registration.get("team_name"):
        data.insert(1, ["Team Squad Name", registration.get("team_name")])

    # Team members if present
    members = registration.get("team_members", [])
    if members:
        member_names = ", ".join(m.get("name", "") for m in members if isinstance(m, dict) and m.get("name"))
        if member_names:
            data.append(["Team Members", member_names])
    
    data.append(["Registration Date", format_value(registration.get("created_at", ""))])
    
    table = Table(data, colWidths=[5.2*cm, 12.2*cm])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#334155")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f8fafc")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 16))

    # Important Guidelines Box
    guide_title = Paragraph("<b>OFFICIAL ENTRY INSTRUCTIONS & GUIDELINES:</b>", ParagraphStyle("GTitle", parent=styles["Normal"], fontSize=9, fontName="Helvetica-Bold", textColor=colors.HexColor("#0f172a"), spaceAfter=5))
    guide_p1 = Paragraph("1. <b>Entry Verification:</b> Present this official PDF pass (on your smartphone or printed) at the campus security and registration verification desk.", ParagraphStyle("G1", parent=styles["Normal"], fontSize=8.5, leading=12, textColor=colors.HexColor("#475569")))
    guide_p2 = Paragraph("2. <b>Identification:</b> Must be accompanied by a valid College Student Photo ID Card for all participants.", ParagraphStyle("G2", parent=styles["Normal"], fontSize=8.5, leading=12, textColor=colors.HexColor("#475569")))
    guide_p3 = Paragraph("3. <b>Schedule & Venue:</b> Check the official schedule on the website and report 30 minutes before your scheduled slot.", ParagraphStyle("G3", parent=styles["Normal"], fontSize=8.5, leading=12, textColor=colors.HexColor("#475569")))
    guide_box = Table([[guide_title], [guide_p1], [guide_p2], [guide_p3]], colWidths=[17.4*cm])
    guide_box.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f5f9")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(guide_box)
    
    elements.append(Spacer(1, 20))
    
    footer_style = ParagraphStyle(
        "Footer",
        parent=styles["Normal"],
        fontSize=8.5,
        textColor=colors.HexColor("#94a3b8"),
        alignment=1,
    )
    elements.append(Paragraph("Authorized by Stark Industries Festival Command & COLORIDO 2K26 Organizing Committee", footer_style))
    elements.append(Paragraph("This is a digitally certified entry document. 100% Free Entry guaranteed.", footer_style))
    
    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_bulk_pdf(registrations: List[dict]) -> io.BytesIO:
    """Generate PDF containing multiple registrations in a table."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm,
                            leftMargin=1*cm, rightMargin=1*cm)
    styles = getSampleStyleSheet()
    elements = []
    
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=18,
        textColor=colors.HexColor("#e94560"),
        alignment=1,
        spaceAfter=10,
    )
    elements.append(Paragraph("COLORIDO 2K26 — Registration Report", title_style))
    elements.append(Spacer(1, 10))
    
    date_style = ParagraphStyle("Date", parent=styles["Normal"], fontSize=9, alignment=1, spaceAfter=15)
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total: {len(registrations)}", date_style))
    
    # Table headers
    headers = ["Reg ID", "Name", "College", "Event", "Phone", "Status"]
    data = [headers]
    
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7, leading=9)
    
    for reg in registrations:
        row = [
            Paragraph(reg.get("registration_id", ""), cell_style),
            Paragraph(reg.get("participant_name", ""), cell_style),
            Paragraph(reg.get("college", ""), cell_style),
            Paragraph(reg.get("event_name", ""), cell_style),
            Paragraph(reg.get("phone", ""), cell_style),
            Paragraph(reg.get("status", ""), cell_style),
        ]
        data.append(row)
    
    table = Table(data, colWidths=[3*cm, 3.5*cm, 4*cm, 3.5*cm, 2.5*cm, 2*cm])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5fa")]),
    ]))
    elements.append(table)
    
    doc.build(elements)
    buffer.seek(0)
    return buffer
